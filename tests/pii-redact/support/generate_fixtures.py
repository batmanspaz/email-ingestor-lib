#!/usr/bin/env python3
"""Builds SYNTHETIC PII-redaction fixtures with PyMuPDF. No real data, ever.

Usage: generate_fixtures.py <outDir>
Writes the fixture files plus <outDir>/manifest.json describing, per fixture, which synthetic
secrets it contains (keys of values.json 'secrets') and which negative controls it carries.
This script is part of the P0 test harness (GREEN by design): the RED tests need it to work so
they fail on the engine, never on the fixtures.
"""
import json
import os
import sys
import zipfile

import fitz  # PyMuPDF

HERE = os.path.dirname(os.path.abspath(__file__))
V = json.load(open(os.path.join(HERE, "values.json")))
S = {k: v["text"] for k, v in V["secrets"].items()}
N = V["negatives"]

FULL_SECRETS = ["ssn1", "ssn2", "itin", "routing", "account", "dob", "pin", "agi", "dl"]


def body_lines():
    return [
        "Form 1040 - U.S. Individual Income Tax Return (SYNTHETIC TEST DOCUMENT)",
        f"Taxpayer: JANE Q TESTER    SSN: {S['ssn1']}",
        f"Spouse: JOHN TESTER    SSN: {S['ssn2']}",
        f"Dependent (ITIN): {S['itin']}",
        f"Date of birth: {S['dob']}",
        f"Routing number: {S['routing']}    Account number: {S['account']}",
        f"Self-select PIN: {S['pin']}    Prior year AGI: ${S['agi']}",
        f"Driver license no: {S['dl']} (CA)",
        f"{N['invoice']['label']} {N['invoice']['text']} is not an SSN",
        f"{N['phone']['label']} {N['phone']['text']}    {N['ein']['label']} {N['ein']['text']}",
        f"{N['amount']['label']} {N['amount']['text']}    {N['filing_date']['label']} {N['filing_date']['text']}",
        f"{N['current_agi']['label']} {N['current_agi']['text']}",
    ]


def write_lines(page, lines, y0=72, color=(0, 0, 0)):
    y = y0
    for line in lines:
        page.insert_text((54, y), line, fontsize=10, fontname="helv", color=color)
        y += 16
    return y


def neg_keys():
    return list(N.keys())


def make_text_layer(path):
    doc = fitz.open()
    p = doc.new_page()
    write_lines(p, body_lines())
    doc.new_page()
    write_lines(doc[1], ["Page 2: continuation", f"Taxpayer SSN again: {S['ssn1']}"])
    doc.save(path)


def image_only_from(text_pdf_path, out_path, pages=None):
    src = fitz.open(text_pdf_path)
    doc = fitz.open()
    for i, pg in enumerate(src):
        if pages is not None and i not in pages:
            continue
        pix = pg.get_pixmap(dpi=200)
        np = doc.new_page(width=pg.rect.width, height=pg.rect.height)
        np.insert_image(np.rect, stream=pix.tobytes("png"))
    doc.save(out_path, deflate=True, garbage=3)


def make_mixed(out, text_pdf):
    src = fitz.open(text_pdf)
    doc = fitz.open()
    doc.insert_pdf(src, from_page=0, to_page=0)
    pix = src[1].get_pixmap(dpi=200)
    np = doc.new_page(width=src[1].rect.width, height=src[1].rect.height)
    np.insert_image(np.rect, stream=pix.tobytes("png"))
    doc.save(out, deflate=True, garbage=3)


def make_acroform(path):
    doc = fitz.open()
    p = doc.new_page()
    write_lines(p, ["Synthetic e-file authorization form"])
    w = fitz.Widget()
    w.field_type = fitz.PDF_WIDGET_TYPE_TEXT
    w.field_name = "taxpayer_ssn"
    w.field_value = S["ssn1"]
    w.rect = fitz.Rect(54, 100, 260, 120)
    p.add_widget(w)
    doc.save(path)


def make_metadata(path):
    doc = fitz.open()
    p = doc.new_page()
    write_lines(p, ["Synthetic document whose body has no private number"])
    doc.set_metadata({
        "title": f"Return for SSN {S['ssn1']}",
        "author": "Synthetic Preparer",
        "subject": f"Taxpayer {S['ssn2']}",
        "keywords": f"dob {S['dob']}",
    })
    doc.set_xml_metadata(
        '<?xpacket begin="" id="W5M0MpCehiHzreSzNTczkc9d"?>'
        f'<x:xmpmeta xmlns:x="adobe:ns:meta/"><rdf:RDF xmlns:rdf="http://www.w3.org/1999/02/22-rdf-syntax-ns#">'
        f'<rdf:Description xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:description>SSN {S["ssn1"]}</dc:description>'
        "</rdf:Description></rdf:RDF></x:xmpmeta><?xpacket end=\"w\"?>"
    )
    doc.save(path)


def make_embedded(path):
    doc = fitz.open()
    p = doc.new_page()
    write_lines(p, ["Synthetic PDF with an embedded worksheet"])
    doc.embfile_add("worksheet.txt", f"Worksheet SSN {S['ssn1']}\n".encode())
    doc.save(path)


def make_annots_toc(path):
    doc = fitz.open()
    p = doc.new_page()
    write_lines(p, ["Synthetic PDF with annotations and a bookmark"])
    p.add_text_annot((72, 200), f"Note: SSN {S['ssn1']}")
    p.add_freetext_annot(fitz.Rect(54, 220, 360, 250), f"Spouse SSN {S['ssn2']}", fontsize=10)
    doc.set_toc([[1, f"Taxpayer {S['ssn1']}", 1]])
    doc.save(path)


def make_hidden_ocg(path):
    doc = fitz.open()
    p = doc.new_page()
    write_lines(p, ["Synthetic PDF with text in a hidden layer"])
    ocg = doc.add_ocg("hidden-layer", on=False)
    p.insert_text((54, 300), f"Hidden SSN {S['ssn1']}", fontsize=10, fontname="helv", oc=ocg)
    doc.save(path)


def make_white_on_white(path):
    doc = fitz.open()
    p = doc.new_page()
    write_lines(p, ["Synthetic PDF with white text"])
    p.insert_text((54, 300), f"Invisible SSN {S['ssn1']}", fontsize=10, fontname="helv", color=(1, 1, 1))
    doc.save(path)


def make_footer_every_page(path, pages=5):
    doc = fitz.open()
    for i in range(pages):
        p = doc.new_page()
        write_lines(p, [f"Synthetic page {i + 1} of {pages}"])
        p.insert_text((54, 800), f"Taxpayer SSN {S['ssn1']}    page {i + 1}", fontsize=8, fontname="helv")
    doc.save(path)


def make_split_spans(path):
    """R2 (Opus D21/H5): the round-1 fixture extracted as ONE span. Now the SSN really is split:
    ssn1 = two spans in DIFFERENT fonts on one baseline; ssn2 = wrapped across two lines;
    itin = two separate text blocks (two-column layout)."""
    doc = fitz.open()
    p = doc.new_page()
    p.insert_text((54, 100), "SSN: 123-45-", fontsize=10, fontname="helv")
    w = fitz.get_text_length("SSN: 123-45-", fontname="helv", fontsize=10)
    p.insert_text((54 + w, 100), "6789", fontsize=10, fontname="cour")
    p.insert_text((54, 140), "Spouse SSN 987-65-", fontsize=10, fontname="helv")
    p.insert_text((54, 154), "4320 (continued)", fontsize=10, fontname="helv")
    tw = fitz.TextWriter(p.rect)
    tw.append((54, 200), "Dependent ITIN 912-78-", fontsize=10)
    tw.write_text(p)
    tw2 = fitz.TextWriter(p.rect)
    tw2.append((320, 200), "1234", fontsize=10)
    tw2.write_text(p)
    doc.save(path)


UNICODE_FONT = "/System/Library/Fonts/Supplemental/Arial Unicode.ttf"


def make_unicode_dashes(path):
    """R2 (Opus D21/H5): base-14 Helvetica cannot encode U+2013/U+00A0/U+2011 (the round-1 en dash extracted as
    U+00B7 and the 'NBSP' line was plain spaces). Embed a Unicode TTF (subset) so the code points really exist."""
    doc = fitz.open()
    p = doc.new_page()
    if os.path.exists(UNICODE_FONT):
        p.insert_font(fontname="UF", fontfile=UNICODE_FONT)
    else:  # non-mac fallback: built-in Droid fallback covers U+2013 at least
        p.insert_font(fontname="UF", fontbuffer=fitz.Font("cjk").buffer)
    p.insert_text((54, 100), "SSN 123 45 6789", fontsize=10, fontname="UF")
    p.insert_text((54, 120), "SSN 987–65–4320", fontsize=10, fontname="UF")
    p.insert_text((54, 140), "ITIN 912‑78‑1234", fontsize=10, fontname="UF")
    doc.subset_fonts()
    doc.save(path, garbage=3, deflate=True)


def make_clean(path):
    doc = fitz.open()
    p = doc.new_page()
    write_lines(p, [
        "Synthetic clean document - no private numbers",
        f"{N['invoice']['label']} {N['invoice']['text']}",
        f"{N['phone']['label']} {N['phone']['text']}   {N['ein']['label']} {N['ein']['text']}",
        f"{N['amount']['label']} {N['amount']['text']}   {N['filing_date']['label']} {N['filing_date']['text']}",
    ])
    doc.save(path)


def make_encrypted(src, path):
    doc = fitz.open(src)
    doc.save(path, encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw="owner-pw", user_pw="user-pw")


def make_corrupt(src, path):
    data = open(src, "rb").read()
    cut = int(len(data) * 0.55)
    open(path, "wb").write(data[:cut] + b"\x00garbage\x00" * 8)


def docx_bytes(paragraphs):
    ct = ('<?xml version="1.0" encoding="UTF-8"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
          '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
          '<Default Extension="xml" ContentType="application/xml"/>'
          '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
    rels = ('<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
    paras = "".join(f"<w:p><w:r><w:t xml:space=\"preserve\">{t}</w:t></w:r></w:p>" for t in paragraphs)
    doc = ('<?xml version="1.0" encoding="UTF-8"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
           f"<w:body>{paras}</w:body></w:document>")
    import io
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", ct)
        z.writestr("_rels/.rels", rels)
        z.writestr("word/document.xml", doc)
    return buf.getvalue()


# ---------------------------------------------------------------------------------------------------------------
# ROUND 2 fixtures (red-team fold-in, 2026-10-14). Every one of these defeats a naive engine/inspector in a way the
# Sonnet/Opus red-team verdicts demonstrated. IDs in brackets map to review-verdicts/{sonnet,opus}-redact-p0.md.
# ---------------------------------------------------------------------------------------------------------------

def _ssn_page(doc, lines=None):
    p = doc.new_page()
    write_lines(p, lines or ["Synthetic return", f"Taxpayer SSN {S['ssn1']}"])
    return p


def make_incremental_history(path):
    """[Opus D1/D2/C2, Sonnet C3] SSN exists ONLY in a superseded revision: rev 1 has it; rev 2 (incremental save)
    redacts the page. Visible text is clean; raw file bytes are not."""
    doc = fitz.open()
    _ssn_page(doc)
    doc.save(path)
    doc.close()
    doc = fitz.open(path)
    pg = doc[0]
    for r in pg.search_for(S["ssn1"]):
        pg.add_redact_annot(r)
    pg.apply_redactions()
    doc.saveIncr()
    doc.close()


def make_orphan_after_redact(path):
    """[Opus D1] apply_redactions + default save (garbage=0) leaves the old content stream as an unreferenced
    object in the file. Visible text clean; raw bytes still hold the SSN."""
    doc = fitz.open()
    _ssn_page(doc)
    doc = fitz.open("pdf", doc.tobytes())
    pg = doc[0]
    for r in pg.search_for(S["ssn1"]):
        pg.add_redact_annot(r)
    pg.apply_redactions()
    doc.save(path)


def make_objstm(path):
    """[Sonnet C3, Opus C1] SSN in a private catalog key, inside a compressed object stream (ObjStm)."""
    doc = fitz.open()
    p = doc.new_page()
    write_lines(p, ["Synthetic PDF with private application data"])
    doc.xref_set_key(doc.pdf_catalog(), "SynthPrivateData", f"(Taxpayer SSN {S['ssn1']})")
    doc.save(path, garbage=1, deflate=True, use_objstms=1)


def make_form_xobject(path):
    """[Sonnet C3] SSN drawn by a Form XObject: not in the page content stream itself."""
    src = fitz.open()
    _ssn_page(src)
    doc = fitz.open()
    p = doc.new_page()
    write_lines(p, ["Synthetic page that places another page as a form XObject"], y0=40)
    p.show_pdf_page(fitz.Rect(36, 80, 576, 780), src, 0)
    doc.save(path, garbage=3, deflate=True)


def make_acroform_full(path):
    """[Opus D6/H1] widget value = ssn1 on a HIDDEN widget (appearance stream carries it), /DV default value = ssn2.
    Blanking the value + update() still leaks; the field must be DELETED."""
    doc = fitz.open()
    p = doc.new_page()
    write_lines(p, ["Synthetic e-file authorization form (hidden field)"])
    w = fitz.Widget()
    w.field_type = fitz.PDF_WIDGET_TYPE_TEXT
    w.field_name = "taxpayer_ssn_hidden"
    w.field_value = S["ssn1"]
    w.rect = fitz.Rect(54, 100, 260, 120)
    p.add_widget(w)
    w2 = list(p.widgets())[0]
    doc.xref_set_key(w2.xref, "DV", f"({S['ssn2']})")
    doc.xref_set_key(w2.xref, "F", "2")  # annotation flag: Hidden
    doc.save(path)


def make_image_on_text_page(path):
    """[Opus D6/C5] a TEXT page (has a text layer, so a naive engine never OCRs it) carrying a scanned stamp image
    with the SSN in its pixels."""
    stamp = fitz.open()
    sp = stamp.new_page(width=300, height=40)
    sp.insert_text((8, 26), f"SSN {S['ssn1']}", fontsize=16, fontname="helv")
    pix = sp.get_pixmap(dpi=200)
    doc = fitz.open()
    p = doc.new_page()
    write_lines(p, ["Synthetic cover letter with a pasted W-2 stamp", "Nothing private is typed on this page."])
    p.insert_image(fitz.Rect(54, 140, 354, 180), stream=pix.tobytes("png"))
    doc.save(path, garbage=3, deflate=True)


# Comb field geometry (PDF points) — the generator records ground truth for the pixel check. [Opus D3/C3]
COMB_X0, COMB_Y0, COMB_W, COMB_H, COMB_GAP = 60, 120, 22, 26, 4


def comb_digit_rects():
    rects = []
    x = COMB_X0
    for i in range(9):
        if i in (3, 5):
            x += 12  # visual dash gap
        rects.append([x + 3, COMB_Y0 + 3, x + COMB_W - 3, COMB_Y0 + COMB_H - 3])
        x += COMB_W + COMB_GAP
    return rects


def make_comb_ssn_image(path):
    """[Opus D3/C3] 1040/W-2 style comb field: each SSN digit sits in its own box under the label. Rasterised to an
    image-only page. Tesseract misreads these (probe: 's|7|8]o'), Claude vision does not."""
    src = fitz.open()
    p = src.new_page()
    p.insert_text((60, 100), "Your social security number", fontsize=9, fontname="helv")
    digits = S["ssn1"].replace("-", "")
    x = COMB_X0
    for i, ch in enumerate(digits):
        if i in (3, 5):
            x += 12
        p.draw_rect(fitz.Rect(x, COMB_Y0, x + COMB_W, COMB_Y0 + COMB_H), color=(0, 0, 0), width=1.2)
        p.insert_text((x + 6, COMB_Y0 + 19), ch, fontsize=14, fontname="cour")
        x += COMB_W + COMB_GAP
    p.insert_text((60, 200), "Synthetic form - comb field test", fontsize=9, fontname="helv")
    image_pages_from(src, path, dpi=200)


def image_pages_from(src, out_path, dpi=200, rotate=0.0, gray=False):
    doc = fitz.open()
    for pg in src:
        pix = pg.get_pixmap(dpi=dpi, colorspace=fitz.csGRAY if gray else fitz.csRGB)
        png = pix.tobytes("png")
        if rotate:
            from PIL import Image
            import io
            im = Image.open(io.BytesIO(png)).convert("RGB").rotate(rotate, expand=False, fillcolor=(255, 255, 255))
            b = io.BytesIO()
            im.save(b, "PNG")
            png = b.getvalue()
        np_ = doc.new_page(width=pg.rect.width, height=pg.rect.height)
        np_.insert_image(np_.rect, stream=png)
    doc.save(out_path, deflate=True, garbage=3)


def make_scan_variants(path):
    """[Sonnet C3] image-only pages that stress local OCR: 3-degree skew, 100 dpi, light-grey ink, /Rotate 90."""
    def page_with(color):
        s = fitz.open()
        pg = s.new_page()
        write_lines(pg, ["Synthetic scan", f"Taxpayer SSN {S['ssn1']}"], color=color)
        return s
    tmp = []
    for i, (dpi, rot, col) in enumerate([(200, 3.0, (0, 0, 0)), (100, 0.0, (0, 0, 0)), (200, 0.0, (0.72, 0.72, 0.72))]):
        t = path + f".part{i}.pdf"
        image_pages_from(page_with(col), t, dpi=dpi, rotate=rot)
        tmp.append(t)
    doc = fitz.open()
    for t in tmp:
        doc.insert_pdf(fitz.open(t))
        os.remove(t)
    rp = page_with((0, 0, 0))
    image_pages_from(rp, path + ".rot.pdf", dpi=200)
    doc.insert_pdf(fitz.open(path + ".rot.pdf"))
    os.remove(path + ".rot.pdf")
    doc[-1].set_rotation(90)
    doc.save(path, garbage=3, deflate=True)


def make_js_action(path):
    """[Opus H2] SSN inside a document-level JavaScript OpenAction."""
    doc = fitz.open()
    p = doc.new_page()
    write_lines(p, ["Synthetic PDF with a JavaScript action"])
    doc.xref_set_key(doc.pdf_catalog(), "OpenAction", f"<</S/JavaScript/JS(var t='{S['ssn1']}';)>>")
    doc.save(path)


def make_xfa(path):
    """[Sonnet C3, Opus H2] XFA datasets stream carrying the SSN (AcroForm /XFA)."""
    doc = fitz.open()
    p = doc.new_page()
    write_lines(p, ["Synthetic XFA form"])
    w = fitz.Widget()
    w.field_type = fitz.PDF_WIDGET_TYPE_TEXT
    w.field_name = "note"
    w.field_value = "n/a"
    w.rect = fitz.Rect(54, 100, 200, 120)
    p.add_widget(w)
    x = doc.get_new_xref()
    doc.update_object(x, "<<>>")
    doc.update_stream(x, f"<xfa:datasets xmlns:xfa=\"http://www.xfa.org/schema/xfa-data/1.0/\"><xfa:data><ssn>{S['ssn1']}</ssn></xfa:data></xfa:datasets>".encode())
    af = doc.xref_get_key(doc.pdf_catalog(), "AcroForm")
    if af[0] == "xref":
        doc.xref_set_key(int(af[1].split()[0]), "XFA", f"[(datasets) {x} 0 R]")
    else:
        doc.xref_set_key(doc.pdf_catalog(), "AcroForm/XFA", f"[(datasets) {x} 0 R]")
    doc.save(path)


def make_custom_info_key(path):
    """[Opus H2] a CUSTOM Info key — doc.metadata hides it."""
    doc = fitz.open()
    p = doc.new_page()
    write_lines(p, ["Synthetic PDF with a custom document-info key"])
    doc.set_metadata({"title": "Synthetic", "author": "Synthetic Preparer"})
    info = doc.xref_get_key(-1, "Info")
    xref = int(info[1].split()[0])
    doc.xref_set_key(xref, "TaxpayerSSN", f"({S['ssn1']})")
    doc.save(path)


def make_link_and_labels(path):
    """[Opus H2, Sonnet C3] SSN in a link URI query string, in /PageLabels, and in a FileAttachment annotation."""
    doc = fitz.open()
    p = doc.new_page()
    write_lines(p, ["Synthetic PDF with a link, page labels and a file-attachment annotation"])
    p.insert_link({"kind": fitz.LINK_URI, "from": fitz.Rect(54, 100, 200, 120), "uri": f"https://example.test/r?ssn={S['ssn1'].replace('-', '')}"})
    doc.set_page_labels([{"startpage": 0, "prefix": f"SSN {S['ssn2']} ", "style": "D", "firstpagenum": 1}])
    p.add_file_annot(fitz.Point(300, 300), f"Worksheet ITIN {S['itin']}\n".encode(), "worksheet.txt")
    doc.save(path)


def make_encrypted_owner_only(src, path):
    """[Sonnet C3] encrypted with an EMPTY user password: opens without a password but is encrypted."""
    doc = fitz.open(src)
    doc.save(path, encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw="owner-only", user_pw="")


def docx_full_bytes():
    """[Sonnet H1, Opus H6] every OOXML part that can carry text, incl. split runs and a tracked deletion."""
    import io
    W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    ct = ('<?xml version="1.0" encoding="UTF-8"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
          '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
          '<Default Extension="xml" ContentType="application/xml"/>'
          '<Default Extension="xlsx" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"/>'
          '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
    rels = ('<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
    doc = (f'<?xml version="1.0" encoding="UTF-8"?><w:document {W}><w:body>'
           '<w:p><w:r><w:t xml:space="preserve">Taxpayer SSN 123-45-</w:t></w:r><w:r><w:rPr><w:b/></w:rPr><w:t>6789</w:t></w:r></w:p>'
           f'<w:p><w:del w:id="1" w:author="x"><w:r><w:delText>Spouse SSN {S["ssn2"]}</w:delText></w:r></w:del></w:p>'
           f'<w:p><w:r><w:t xml:space="preserve">{N["ein"]["label"]} {N["ein"]["text"]} {N["current_agi"]["label"]} {N["current_agi"]["text"]}</w:t></w:r></w:p>'
           '</w:body></w:document>')
    part = lambda tag, body: f'<?xml version="1.0" encoding="UTF-8"?><w:{tag} {W}><w:p><w:r><w:t xml:space="preserve">{body}</w:t></w:r></w:p></w:{tag}>'
    core = ('<?xml version="1.0" encoding="UTF-8"?><cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
            f'xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>Return for SSN {S["ssn1"]}</dc:title></cp:coreProperties>')
    custom = ('<?xml version="1.0" encoding="UTF-8"?><Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties" '
              f'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"><property name="DL"><vt:lpwstr>Driver license no: {S["dl"]}</vt:lpwstr></property></Properties>')
    cxml = f'<?xml version="1.0" encoding="UTF-8"?><data><agi>Prior year AGI: ${S["agi"]}</agi></data>'
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", ct)
        z.writestr("_rels/.rels", rels)
        z.writestr("word/document.xml", doc)
        z.writestr("word/comments.xml", part("comments", f"Dependent ITIN {S['itin']}"))
        z.writestr("word/header1.xml", part("hdr", f"Date of birth: {S['dob']}"))
        z.writestr("word/footer1.xml", part("ftr", f"Routing number: {S['routing']} Account number: {S['account']}"))
        z.writestr("word/footnotes.xml", part("footnotes", f"Self-select PIN: {S['pin']}"))
        z.writestr("docProps/core.xml", core)
        z.writestr("docProps/custom.xml", custom)
        z.writestr("customXml/item1.xml", cxml)
        z.writestr("word/embeddings/Microsoft_Excel_Worksheet.xlsx", xlsx_bytes(minimal=True))
    return buf.getvalue()


def xlsx_bytes(minimal=False):
    """[Sonnet H1, Opus H7] SSN as a NUMERIC cell with numFmt 000-00-0000, shared strings, a hidden sheet, a cell
    comment, a defined name and a document title."""
    import io
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.workbook.defined_name import DefinedName
    wb = Workbook()
    ws = wb.active
    ws.title = "Return"
    ws["A1"] = "SSN"
    ws["A2"] = int(S["ssn1"].replace("-", ""))
    ws["A2"].number_format = "000-00-0000"
    if not minimal:
        ws["B1"] = "Routing number"
        ws["B2"] = S["routing"]
        ws["C1"] = "Account number"
        ws["C2"] = S["account"]
        ws["A2"].comment = Comment(f"Dependent ITIN {S['itin']}", "synthetic")
        hs = wb.create_sheet("Hidden")
        hs["A1"] = f"Spouse SSN {S['ssn2']}"
        hs.sheet_state = "hidden"
        dn = DefinedName("TaxpayerDOB", attr_text=f'"Date of birth: {S["dob"]}"')
        try:
            wb.defined_names["TaxpayerDOB"] = dn
        except TypeError:
            wb.defined_names.append(dn)
        wb.properties.title = f"Account number: {S['account']}"
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def rtf_bytes():
    return ("{\\rtf1\\ansi\\deff0{\\fonttbl{\\f0 Helvetica;}}\\f0\\fs20 Synthetic RTF\\par "
            f"Taxpayer SSN {S['ssn1']}\\par Date of birth: {S['dob']}\\par }}").encode("latin-1")


def eml_bytes():
    """[Sonnet H2] a saved email: SSN in a header AND in a base64 body part (no ASCII SSN in the raw bytes)."""
    import base64
    body = base64.b64encode(f"Taxpayer SSN {S['ssn2']}\r\n".encode()).decode()
    return ("From: Synthetic Preparer <preparer@example.test>\r\nTo: tester@example.test\r\n"
            f"Subject: Your return {S['ssn1']}\r\nMIME-Version: 1.0\r\nContent-Type: text/plain; charset=utf-8\r\n"
            f"Content-Transfer-Encoding: base64\r\n\r\n{body}\r\n").encode()


def legacy_binary_with_utf16(label):
    """[Sonnet H2] a legacy binary office file (.doc/.xls) whose text is UTF-16LE: a latin1 raw-byte scan sees no SSN."""
    return b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 504 + f"{label} SSN {S['ssn1']}".encode("utf-16-le") + b"\x00" * 64


def make_photos(out, text_pdf):
    """[Opus D4/C4] photographed / scanned image FILES (the producer sends these straight to Claude vision today)."""
    from PIL import Image
    import io
    src = fitz.open(text_pdf)
    pg = src[0]
    pix = pg.get_pixmap(dpi=150)
    im = Image.open(io.BytesIO(pix.tobytes("png"))).convert("RGB")
    truth = []
    for key in ("ssn1", "ssn2", "itin"):
        for r in pg.search_for(S[key]):
            s = 150 / 72
            # first 5 digits (area + group) of the SSN only; last 4 may legitimately stay visible
            w5 = fitz.get_text_length(S[key][:6], fontname="helv", fontsize=10)
            truth.append([int(r.x0 * s), int(r.y0 * s), int((r.x0 + w5) * s), int(r.y1 * s)])
    im.save(os.path.join(out, "w2_scan.png"))
    im.save(os.path.join(out, "w2.tiff"))
    exif = Image.Exif()
    exif[0x010E] = f"Scan of SSN {S['ssn1']}"  # ImageDescription
    im.rotate(2, expand=False, fillcolor=(255, 255, 255)).save(os.path.join(out, "w2_photo.jpg"), "JPEG", quality=60, exif=exif)
    heic = os.path.join(out, "w2.heic")
    try:
        import subprocess
        subprocess.run(["sips", "-s", "format", "heic", os.path.join(out, "w2_scan.png"), "--out", heic], check=True, capture_output=True)
    except Exception:
        pass
    return truth, os.path.exists(heic)


def make_zip_variants(P):
    """[Sonnet C2, Opus H10] SSN in member NAMES, nested zip, encrypted member, zip-slip member, zip bomb."""
    with zipfile.ZipFile(P("zip_member_names.zip"), "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr(f"{S['ssn1']}/notes.txt", "nothing private inside\n")
        z.writestr(f"W2 {S['ssn2']}.txt", "nothing private inside either\n")
    with zipfile.ZipFile(P("nested.zip"), "w", zipfile.ZIP_DEFLATED) as z:
        z.write(P("bundle.zip"), "inner/bundle.zip")
    with zipfile.ZipFile(P("zip_slip.zip"), "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("../../escaped.txt", f"Taxpayer SSN {S['ssn1']}\n")
    import subprocess
    tmpd = P("_enc_src")
    os.makedirs(tmpd, exist_ok=True)
    open(os.path.join(tmpd, "secret.txt"), "w").write(f"Taxpayer SSN {S['ssn1']}\n")
    subprocess.run(["zip", "-q", "-j", "-P", "zip-pw", P("zip_encrypted_member.zip"), os.path.join(tmpd, "secret.txt")], check=True)
    import shutil
    shutil.rmtree(tmpd)
    with zipfile.ZipFile(P("zip_bomb.zip"), "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("zeros.txt", b"0" * (5 * 1024 * 1024))


def write_round2(out, P):
    text_pdf = P("text_layer.pdf")
    make_incremental_history(P("incremental_history.pdf"))
    make_orphan_after_redact(P("orphan_after_redact.pdf"))
    make_objstm(P("objstm.pdf"))
    make_form_xobject(P("form_xobject.pdf"))
    make_acroform_full(P("acroform_full.pdf"))
    make_image_on_text_page(P("image_on_text_page.pdf"))
    make_comb_ssn_image(P("comb_ssn_image.pdf"))
    make_scan_variants(P("scan_variants.pdf"))
    make_js_action(P("js_action.pdf"))
    make_xfa(P("xfa.pdf"))
    make_custom_info_key(P("custom_info_key.pdf"))
    make_link_and_labels(P("link_labels_fileannot.pdf"))
    make_encrypted_owner_only(text_pdf, P("encrypted_owner_only.pdf"))
    open(P("return_full.docx"), "wb").write(docx_full_bytes())
    open(P("sheet.xlsx"), "wb").write(xlsx_bytes())
    txt = "\n".join(body_lines()) + "\n"
    open(P("notes_utf16le.txt"), "wb").write(b"\xff\xfe" + txt.encode("utf-16-le"))
    open(P("notes_utf16be.txt"), "wb").write(b"\xfe\xff" + txt.encode("utf-16-be"))
    open(P("notes_utf8bom.txt"), "wb").write(b"\xef\xbb\xbf" + txt.encode("utf-8"))
    open(P("notes_crlf.txt"), "wb").write(txt.replace("\n", "\r\n").encode("utf-8"))
    open(P("sheet_utf16.csv"), "wb").write(b"\xff\xfe" + f"name,ssn,routing\r\nJane Tester,{S['ssn1']},{S['routing']}\r\n".encode("utf-16-le"))
    open(P("page_latin1.html"), "wb").write(f"<html><body><p>Café receipt. Taxpayer SSN: {S['ssn1']}</p></body></html>\n".encode("latin-1"))
    open(P("blob_utf16.bin"), "wb").write(b"\x00\x01\x02\x03" + f"SSN {S['ssn1']}".encode("utf-16-le") + b"\x04\x05")
    rows = "\n".join(f"Tester {i},2345678{i:02d},{S['routing']}" for i in range(50))
    open(P("ssn_column.csv"), "w").write(f"name,SSN,Routing number\n{rows}\n")
    open(P("split_columns.csv"), "w").write('name,ssn_area,ssn_group,ssn_serial\n"Jane Tester","123","45","6789"\n')
    open(P("html_entities.html"), "w").write(
        "<html><body><p>Taxpayer SSN &#49;23-45-6789</p><p>Spouse SSN 987&ndash;65&ndash;4320</p>"
        "<p>ITIN 912-78-<b>1234</b></p></body></html>\n")
    open(P("json_escapes.json"), "w").write('{"ssn": "123\\u002d45\\u002d6789", "spouse": "987-65-4320"}\n')
    open(P("notes.md"), "w").write(f"# Synthetic\n\n- Taxpayer SSN: {S['ssn1']}\n- Date of birth: {S['dob']}\n")
    open(P("notes.xml"), "w").write(f"<r><ssn>{S['ssn1']}</ssn></r>\n")
    open(P("notes.log"), "w").write(f"2026-10-14 processed taxpayer SSN {S['ssn1']}\n")
    open(P("notes.rtf"), "wb").write(rtf_bytes())
    open(P("return.eml"), "wb").write(eml_bytes())
    open(P("legacy.doc"), "wb").write(legacy_binary_with_utf16("doc"))
    open(P("legacy.xls"), "wb").write(legacy_binary_with_utf16("xls"))
    photo_truth, has_heic = make_photos(out, text_pdf)
    make_zip_variants(P)
    return photo_truth, has_heic


def round2_manifest(photo_truth, has_heic):
    T = "text"
    m = {
        "incremental_history.pdf": {"secrets": ["ssn1"], "surfaces": ["filebytes"], "visible_clean": True, "pages": 1},
        "orphan_after_redact.pdf": {"secrets": ["ssn1"], "surfaces": ["filebytes"], "visible_clean": True, "pages": 1},
        "objstm.pdf": {"secrets": ["ssn1"], "surfaces": ["objects"], "visible_clean": True, "pages": 1},
        "form_xobject.pdf": {"secrets": ["ssn1"], "surfaces": ["text", "objects"], "pages": 1},
        "acroform_full.pdf": {"secrets": ["ssn1", "ssn2"], "surfaces": ["widgets", "objects"], "pages": 1},
        "image_on_text_page.pdf": {"secrets": ["ssn1"], "surfaces": ["ocr"], "visible_clean": True, "pages": 1},
        "comb_ssn_image.pdf": {"secrets": ["ssn1"], "surfaces": ["pixels"], "pages": 1,
                               "pixel_truth": [{"page": 0, "rect": r, "dpi": 200} for r in comb_digit_rects()[:5]]},
        "scan_variants.pdf": {"secrets": ["ssn1"], "surfaces": ["ocr"], "pages": 4},
        "js_action.pdf": {"secrets": ["ssn1"], "surfaces": ["objects"], "visible_clean": True, "pages": 1},
        "xfa.pdf": {"secrets": ["ssn1"], "surfaces": ["objects"], "visible_clean": True, "pages": 1},
        "custom_info_key.pdf": {"secrets": ["ssn1"], "surfaces": ["objects"], "visible_clean": True, "pages": 1},
        "link_labels_fileannot.pdf": {"secrets": ["ssn1", "ssn2", "itin"], "surfaces": ["objects"], "visible_clean": True, "pages": 1},
        "encrypted_owner_only.pdf": {"secrets": FULL_SECRETS, "surfaces": ["text"], "pages": 2},
        "return_full.docx": {"secrets": FULL_SECRETS, "negatives": ["ein", "current_agi"], "kind": "docx"},
        "sheet.xlsx": {"secrets": ["ssn1", "ssn2", "itin", "routing", "account", "dob"], "kind": "xlsx"},
        "notes_utf16le.txt": {"secrets": FULL_SECRETS, "kind": T, "encoding": "utf-16"},
        "notes_utf16be.txt": {"secrets": FULL_SECRETS, "kind": T, "encoding": "utf-16"},
        "notes_utf8bom.txt": {"secrets": FULL_SECRETS, "kind": T, "encoding": "utf-8"},
        "notes_crlf.txt": {"secrets": FULL_SECRETS, "kind": T, "encoding": "utf-8"},
        "sheet_utf16.csv": {"secrets": ["ssn1", "routing"], "kind": T, "encoding": "utf-16"},
        "page_latin1.html": {"secrets": ["ssn1"], "kind": T, "encoding": "latin1"},
        "blob_utf16.bin": {"secrets": ["ssn1"], "kind": "binary", "expect": "held:unsupported"},
        "ssn_column.csv": {"secrets": [], "kind": T, "column_ssn": {"column": 1, "rows": 50}},
        "split_columns.csv": {"secrets": ["ssn1"], "kind": T},
        "html_entities.html": {"secrets": ["ssn1", "ssn2", "itin"], "kind": "html"},
        "json_escapes.json": {"secrets": ["ssn1", "ssn2"], "kind": "json"},
        "notes.md": {"secrets": ["ssn1", "dob"], "kind": T},
        "notes.xml": {"secrets": ["ssn1"], "kind": T},
        "notes.log": {"secrets": ["ssn1"], "kind": T},
        "notes.rtf": {"secrets": ["ssn1", "dob"], "kind": "rtf"},
        "return.eml": {"secrets": ["ssn1", "ssn2"], "kind": "eml"},
        "legacy.doc": {"secrets": ["ssn1"], "kind": "binary"},
        "legacy.xls": {"secrets": ["ssn1"], "kind": "binary"},
        "w2_scan.png": {"secrets": FULL_SECRETS, "kind": "image", "pixel_truth": [{"rect": r} for r in photo_truth]},
        "w2.tiff": {"secrets": FULL_SECRETS, "kind": "image", "pixel_truth": [{"rect": r} for r in photo_truth]},
        "w2_photo.jpg": {"secrets": FULL_SECRETS, "kind": "image", "exif": True},
        "zip_member_names.zip": {"secrets": ["ssn1", "ssn2"], "kind": "zip", "names_only": True},
        "nested.zip": {"secrets": FULL_SECRETS, "kind": "zip"},
        "zip_slip.zip": {"secrets": ["ssn1"], "kind": "zip"},
        "zip_encrypted_member.zip": {"secrets": ["ssn1"], "kind": "zip", "expect": "held:encrypted"},
        "zip_bomb.zip": {"secrets": [], "kind": "zip", "expect": "held:too_large"},
    }
    if has_heic:
        m["w2.heic"] = {"secrets": FULL_SECRETS, "kind": "image"}
    return m


# ---------------------------------------------------------------------------------------------------------------
# ROUND 3 fixtures (Fable red-team, review-verdicts/fable-redact-p0-r2.md). Pixels that a PAGE RENDER never shows.
# ---------------------------------------------------------------------------------------------------------------
SSN_KEYS_ON_PAGE0 = ("ssn1", "ssn2", "itin")


def first5_rects_pt(page):
    """First-5-digit rects (PDF points) of every SSN-class secret on a text_layer page."""
    out = []
    for key in SSN_KEYS_ON_PAGE0:
        for r in page.search_for(S[key]):
            w5 = fitz.get_text_length(S[key][:6], fontname="helv", fontsize=10)
            out.append(fitz.Rect(r.x0, r.y0, r.x0 + w5, r.y1))
    return out


def _stamp_png(text=None, dpi=200, w=300, h=40, size=16):
    st = fitz.open()
    sp = st.new_page(width=w, height=h)
    sp.insert_text((8, 26), text or f"SSN {S['ssn1']}", fontsize=size, fontname="helv")
    return sp.get_pixmap(dpi=dpi).tobytes("png")


def make_overlay_rect_image(text_pdf, image_pdf, path):
    """[Fable C1 / naive #7] image-only page with FILLED RECTANGLES drawn over the SSNs; the image XObject pixels are
    untouched. Page render looks redacted; the XObject at native resolution is not."""
    src = fitz.open(text_pdf)
    rects = first5_rects_pt(src[0])
    doc = fitz.open(image_pdf)
    for i in range(len(doc)):
        for key in SSN_KEYS_ON_PAGE0:
            for r in src[i].search_for(S[key]):
                doc[i].draw_rect(r, color=(0, 0, 0), fill=(0, 0, 0), overlay=True)
    doc.save(path, garbage=4, deflate=True, clean=True)
    return rects


def make_thumb(text_pdf, path):
    """[Fable C1 / naive #8] page text properly redacted, but a /Thumb page thumbnail rendered BEFORE redaction stays."""
    pix = fitz.open(text_pdf)[0].get_pixmap(dpi=150)
    d = fitz.open(text_pdf)
    for pg in d:
        for key in SSN_KEYS_ON_PAGE0:
            for r in pg.search_for(S[key]):
                pg.add_redact_annot(r, fill=(0, 0, 0))
        pg.apply_redactions()
    x = d.get_new_xref()
    d.update_object(x, f"<</Type/XObject/Subtype/Image/Width {pix.width}/Height {pix.height}/ColorSpace/DeviceRGB/BitsPerComponent 8>>")
    d.update_stream(x, pix.samples)
    d.xref_set_key(d[0].xref, "Thumb", f"{x} 0 R")
    d.save(path, garbage=4, deflate=True, clean=True)


def make_hidden_ocg_image(path):
    """[Fable C1 / naive #9] an image in a HIDDEN optional-content layer: the page render honours visibility."""
    d = fitz.open()
    p = d.new_page()
    write_lines(p, ["Synthetic cover page, nothing typed here"])
    ocg = d.add_ocg("hidden", on=False)
    p.insert_image(fitz.Rect(54, 140, 354, 180), stream=_stamp_png(), oc=ocg)
    d.save(path, garbage=4, deflate=True, clean=True)


def make_tiny_image(path):
    """[Fable C1 / naive #10] a 30x4 pt pasted image that is 1250x167 px natively: unreadable in a 220/300 dpi page
    render, readable at native resolution (and by a vision model)."""
    d = fitz.open()
    p = d.new_page()
    write_lines(p, ["Synthetic page with a very small pasted image"])
    p.insert_image(fitz.Rect(54, 100, 84, 104), stream=_stamp_png(dpi=300))
    d.save(path, garbage=4, deflate=True, clean=True)


def make_smask_image(path):
    """[Fable C1] the SSN lives only in an image's SOFT MASK (alpha): white RGB on a white page renders as nothing."""
    from PIL import Image
    import io
    mask = Image.open(io.BytesIO(_stamp_png())).convert("L").point(lambda v: 255 - v)  # text = opaque
    rgba = Image.new("RGBA", mask.size, (255, 255, 255, 0))
    rgba.putalpha(mask)
    b = io.BytesIO()
    rgba.save(b, "PNG")
    d = fitz.open()
    p = d.new_page()
    write_lines(p, ["Synthetic page with a white-on-white masked image"])
    p.insert_image(fitz.Rect(54, 140, 354, 180), stream=b.getvalue())
    d.save(path, garbage=4, deflate=True, clean=True)


def make_embedded_binaries(P):
    """[Fable C1/M5] embedded files that are NOT text: an image (embfile) and a whole PDF (FileAttachment annot)."""
    d = fitz.open()
    p = d.new_page()
    write_lines(p, ["Synthetic PDF carrying an embedded scan"])
    d.embfile_add("scan.png", _stamp_png())
    d.save(P("embedded_image.pdf"))
    d = fitz.open()
    p = d.new_page()
    write_lines(p, ["Synthetic PDF carrying an attached return"])
    p.add_file_annot(fitz.Point(300, 300), open(P("text_layer.pdf"), "rb").read(), "return.pdf")
    d.save(P("embedded_pdf.pdf"))


def make_scan_skew_rotate(path):
    """[Fable M1] the 'straightforward' scan pages (3-degree skew, /Rotate 90 at 200 dpi): these MUST be redacted."""
    src = fitz.open(path.replace("scan_skew_rotate.pdf", "scan_variants.pdf"))
    d = fitz.open()
    d.insert_pdf(src, from_page=0, to_page=0)
    d.insert_pdf(src, from_page=3, to_page=3)
    d.save(path, garbage=3, deflate=True)


def make_comb_redacted_ok(path):
    """[Fable M4] what a CORRECT comb redaction looks like: label, first five boxes solid-filled, last 4 visible.
    The verifier must accept it (a uniform fill in the value window is a located MASKED value)."""
    src = fitz.open()
    p = src.new_page()
    p.insert_text((60, 100), "Your social security number", fontsize=9, fontname="helv")
    digits = S["ssn1"].replace("-", "")
    x = COMB_X0
    for i, ch in enumerate(digits):
        if i in (3, 5):
            x += 12
        r = fitz.Rect(x, COMB_Y0, x + COMB_W, COMB_Y0 + COMB_H)
        if i < 5:
            p.draw_rect(r, color=(0, 0, 0), fill=(0, 0, 0), width=1.2)
        else:
            p.draw_rect(r, color=(0, 0, 0), width=1.2)
            p.insert_text((x + 6, COMB_Y0 + 19), ch, fontsize=14, fontname="cour")
        x += COMB_W + COMB_GAP
    image_pages_from(src, path, dpi=200)


def _jpeg_with_app1(jpeg_bytes, app1_payload):
    seg = b"\xff\xe1" + (len(app1_payload) + 2).to_bytes(2, "big") + app1_payload
    return jpeg_bytes[:2] + seg + jpeg_bytes[2:]


def _exif_with_thumbnail(thumb_jpeg):
    """Minimal big-endian EXIF: empty IFD0 -> IFD1 {Compression=6, JPEGInterchangeFormat, ...Length} -> thumbnail."""
    import struct
    tiff = b"MM\x00\x2a" + struct.pack(">I", 8)
    ifd0 = struct.pack(">H", 0) + struct.pack(">I", 14)
    thumb_off = 14 + 2 + 3 * 12 + 4
    ifd1 = struct.pack(">H", 3)
    ifd1 += struct.pack(">HHII", 0x0103, 3, 1, 6 << 16)
    ifd1 += struct.pack(">HHII", 0x0201, 4, 1, thumb_off)
    ifd1 += struct.pack(">HHII", 0x0202, 4, 1, len(thumb_jpeg))
    ifd1 += struct.pack(">I", 0)
    return b"Exif\x00\x00" + tiff + ifd0 + ifd1 + thumb_jpeg


def make_image_containers(P, text_pdf):
    """[Fable M2] image files whose secret is in a SECONDARY payload: EXIF text with clean pixels, PNG tEXt/iTXt with
    clean pixels, an EXIF thumbnail of the unredacted scan, and a 2-frame TIFF whose 2nd frame has the SSN."""
    from PIL import Image, PngImagePlugin
    import io
    clean = fitz.open()
    cp = clean.new_page(width=400, height=120)
    cp.insert_text((10, 60), "Grocery receipt total $12.34", fontsize=16, fontname="helv")
    clean_im = Image.open(io.BytesIO(cp.get_pixmap(dpi=150).tobytes("png"))).convert("RGB")
    exif = Image.Exif()
    exif[0x010E] = f"Scan of SSN {S['ssn1']}"
    clean_im.save(P("clean_pixels_exif_ssn.jpg"), "JPEG", quality=85, exif=exif)
    info = PngImagePlugin.PngInfo()
    info.add_text("Comment", f"Taxpayer SSN {S['ssn1']}")
    info.add_itxt("Description", f"Spouse SSN {S['ssn2']}")
    clean_im.save(P("w2_meta.png"), "PNG", pnginfo=info)
    w2 = Image.open(io.BytesIO(fitz.open(text_pdf)[0].get_pixmap(dpi=150).tobytes("png"))).convert("RGB")
    top = w2.crop((0, 0, w2.width, w2.height // 4))
    tb = io.BytesIO()
    top.resize((900, int(900 * top.height / top.width))).save(tb, "JPEG", quality=60)
    main = io.BytesIO()
    w2.save(main, "JPEG", quality=70)
    open(P("w2_thumb.jpg"), "wb").write(_jpeg_with_app1(main.getvalue(), _exif_with_thumbnail(tb.getvalue())))
    Image.frombytes("RGB", clean_im.size, clean_im.tobytes()).save(P("w2_2page.tiff"), save_all=True, append_images=[Image.frombytes("RGB", w2.size, w2.tobytes())], compression="tiff_deflate")


def write_round3(out, P):
    text_pdf, image_pdf = P("text_layer.pdf"), P("image_only.pdf")
    overlay_rects = make_overlay_rect_image(text_pdf, image_pdf, P("overlay_rect_image.pdf"))
    make_thumb(text_pdf, P("thumb.pdf"))
    make_hidden_ocg_image(P("hidden_ocg_image.pdf"))
    make_tiny_image(P("tiny_image.pdf"))
    make_smask_image(P("smask_image.pdf"))
    make_embedded_binaries(P)
    make_scan_skew_rotate(P("scan_skew_rotate.pdf"))
    make_comb_redacted_ok(P("comb_redacted_ok.pdf"))
    make_image_containers(P, text_pdf)
    s = 200 / 72  # image_only.pdf / overlay fixture: page 0 image was rendered at 200 dpi from text_layer page 0
    xtruth = [{"page": 0, "rect": [int(r.x0 * s), int(r.y0 * s), int(r.x1 * s), int(r.y1 * s)]} for r in overlay_rects]
    return xtruth


def round3_manifest(xtruth):
    return {
        "overlay_rect_image.pdf": {"secrets": FULL_SECRETS, "surfaces": ["images"], "visible_clean_ssn": True, "pages": 2, "xobject_truth": xtruth},
        "thumb.pdf": {"secrets": ["ssn1", "ssn2", "itin"], "surfaces": ["images"], "visible_clean_ssn": True, "pages": 2},
        "hidden_ocg_image.pdf": {"secrets": ["ssn1"], "surfaces": ["images"], "visible_clean": True, "pages": 1},
        "tiny_image.pdf": {"secrets": ["ssn1"], "surfaces": ["images"], "visible_clean": True, "pages": 1},
        "smask_image.pdf": {"secrets": ["ssn1"], "surfaces": ["images"], "visible_clean": True, "pages": 1},
        "embedded_image.pdf": {"secrets": ["ssn1"], "surfaces": ["embedded_bin"], "visible_clean": True, "pages": 1},
        "embedded_pdf.pdf": {"secrets": FULL_SECRETS, "surfaces": ["embedded_bin"], "visible_clean": True, "pages": 1},
        "scan_skew_rotate.pdf": {"secrets": ["ssn1"], "surfaces": ["ocr"], "pages": 2},
        "comb_redacted_ok.pdf": {"secrets": [], "surfaces": [], "pages": 1},
        "clean_pixels_exif_ssn.jpg": {"secrets": ["ssn1"], "kind": "image", "where": "exif"},
        "w2_meta.png": {"secrets": ["ssn1", "ssn2"], "kind": "image", "where": "png_text"},
        "w2_thumb.jpg": {"secrets": FULL_SECRETS, "kind": "image", "where": "exif_thumbnail"},
        "w2_2page.tiff": {"secrets": FULL_SECRETS, "kind": "image", "where": "frame_2"},
    }


def main():
    out = sys.argv[1]
    os.makedirs(out, exist_ok=True)
    P = lambda n: os.path.join(out, n)

    make_text_layer(P("text_layer.pdf"))
    image_only_from(P("text_layer.pdf"), P("image_only.pdf"))
    make_mixed(P("mixed.pdf"), P("text_layer.pdf"))
    make_acroform(P("acroform.pdf"))
    make_metadata(P("metadata.pdf"))
    make_embedded(P("embedded.pdf"))
    make_annots_toc(P("annots_toc.pdf"))
    make_hidden_ocg(P("hidden_ocg.pdf"))
    make_white_on_white(P("white_on_white.pdf"))
    make_footer_every_page(P("footer_every_page.pdf"))
    make_split_spans(P("split_spans.pdf"))
    make_unicode_dashes(P("unicode_dashes.pdf"))
    make_clean(P("clean.pdf"))
    make_encrypted(P("text_layer.pdf"), P("encrypted.pdf"))
    make_corrupt(P("text_layer.pdf"), P("corrupt.pdf"))

    txt = "\n".join(body_lines()) + "\n"
    open(P("notes.txt"), "w").write(txt)
    open(P("sheet.csv"), "w").write(f"name,ssn,routing\nJane Tester,{S['ssn1']},{S['routing']}\n")
    open(P("page.html"), "w").write(f"<html><body><p>Taxpayer SSN: {S['ssn1']}</p><p>Date of birth: {S['dob']}</p></body></html>\n")
    open(P("return.docx"), "wb").write(docx_bytes(body_lines()))
    with zipfile.ZipFile(P("bundle.zip"), "w", zipfile.ZIP_DEFLATED) as z:
        z.write(P("text_layer.pdf"), "text_layer.pdf")
        z.write(P("notes.txt"), "notes.txt")
        z.write(P("clean.pdf"), "clean.pdf")
    with zipfile.ZipFile(P("bundle_with_bad_member.zip"), "w", zipfile.ZIP_DEFLATED) as z:
        z.write(P("text_layer.pdf"), "text_layer.pdf")
        z.write(P("corrupt.pdf"), "corrupt.pdf")

    # kinds: what surfaces hold secrets (used by the leak inspector sanity tests)
    manifest = {
        "text_layer.pdf": {"secrets": FULL_SECRETS, "negatives": neg_keys(), "surfaces": ["text"], "pages": 2},
        "image_only.pdf": {"secrets": FULL_SECRETS, "negatives": neg_keys(), "surfaces": ["ocr"], "pages": 2},
        "mixed.pdf": {"secrets": ["ssn1", "ssn2", "itin", "routing", "account", "dob", "pin", "agi", "dl"], "negatives": neg_keys(), "surfaces": ["text", "ocr"], "pages": 2},
        "acroform.pdf": {"secrets": ["ssn1"], "negatives": [], "surfaces": ["widgets"], "pages": 1},
        "metadata.pdf": {"secrets": ["ssn1", "ssn2", "dob"], "negatives": [], "surfaces": ["metadata", "xmp"], "pages": 1},
        "embedded.pdf": {"secrets": ["ssn1"], "negatives": [], "surfaces": ["embedded"], "pages": 1},
        "annots_toc.pdf": {"secrets": ["ssn1", "ssn2"], "negatives": [], "surfaces": ["annots", "toc"], "pages": 1},
        "hidden_ocg.pdf": {"secrets": ["ssn1"], "negatives": [], "surfaces": ["raw"], "pages": 1},
        "white_on_white.pdf": {"secrets": ["ssn1"], "negatives": [], "surfaces": ["text"], "pages": 1},
        "footer_every_page.pdf": {"secrets": ["ssn1"], "negatives": [], "surfaces": ["text"], "pages": 5},
        "split_spans.pdf": {"secrets": ["ssn1", "ssn2", "itin"], "negatives": [], "surfaces": ["text"], "pages": 1},
        "unicode_dashes.pdf": {"secrets": ["ssn1", "ssn2", "itin"], "negatives": [], "surfaces": ["text"], "pages": 1},
        "clean.pdf": {"secrets": [], "negatives": ["invoice", "phone", "ein", "amount", "filing_date"], "surfaces": [], "pages": 1},
        "encrypted.pdf": {"secrets": FULL_SECRETS, "negatives": [], "surfaces": [], "pages": 2, "expect": "held:encrypted"},
        "corrupt.pdf": {"secrets": [], "negatives": [], "surfaces": [], "pages": 0, "expect": "held:corrupt"},
        "notes.txt": {"secrets": FULL_SECRETS, "negatives": neg_keys(), "kind": "text"},
        "sheet.csv": {"secrets": ["ssn1", "routing"], "negatives": [], "kind": "text"},
        "page.html": {"secrets": ["ssn1", "dob"], "negatives": [], "kind": "text"},
        "return.docx": {"secrets": FULL_SECRETS, "negatives": neg_keys(), "kind": "docx"},
        "bundle.zip": {"secrets": FULL_SECRETS, "negatives": neg_keys(), "kind": "zip"},
        "bundle_with_bad_member.zip": {"secrets": FULL_SECRETS, "negatives": [], "kind": "zip", "expect": "held:corrupt_member"},
    }
    photo_truth, has_heic = write_round2(out, P)
    manifest.update(round2_manifest(photo_truth, has_heic))
    xtruth = write_round3(out, P)
    manifest.update(round3_manifest(xtruth))
    manifest["image_only.pdf"]["xobject_truth"] = xtruth
    json.dump(manifest, open(P("manifest.json"), "w"), indent=2)
    print(json.dumps({"ok": True, "dir": out, "files": sorted(manifest)}))


if __name__ == "__main__":
    main()
